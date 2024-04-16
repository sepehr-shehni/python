from docx import Document

# ایجاد یک شیء Document جدید
doc = Document()

# ایجاد یک جدول سه ردیف و سه ستون
table = doc.add_table(rows=3, cols=3)

# مقداردهی به سلول‌ها
for i in range(3):
    for j in range(3):
        cell = table.cell(i, j)
        cell.text = f'Row {i+1}, Col {j+1}'

# ذخیره سند با نام مشخص
doc.save('document_with_table.docx')
